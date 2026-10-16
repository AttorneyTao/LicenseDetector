"""Optional input identifiers must survive every input/output path unchanged."""
import asyncio
import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import api
from core.utils import get_input_metadata, read_input_excel


def input_rows():
    return [
        dict(name='slow', version='1', github_url='https://github.com/example/same',
             serial_number='000123', package_manager='cargo'),
        dict(name='fast', version='2', github_url='https://github.com/example/same',
             serial_number='123456789012345678901234', package_manager='npm'),
        dict(name='error', github_url='https://github.com/example/error',
             serial_number='NA', package_manager='None'),
        dict(name='skip', github_url='https://github.com/example/skip',
             serial_number=None, package_manager=None),
        dict(name='fallback', github_url='https://example.com/source.tar.gz',
             serial_number='批次-A01', package_manager='custom'),
    ]


def excel_bytes(rows):
    stream = io.BytesIO()
    pd.DataFrame(rows).to_excel(stream, index=False)
    return stream.getvalue()


def test_reader_keeps_identifiers_and_existing_column_parsing():
    content = excel_bytes(input_rows())
    original = pd.read_excel(io.BytesIO(content))
    actual = read_input_excel(io.BytesIO(content))
    pd.testing.assert_frame_equal(actual[['name', 'version', 'github_url']],
                                  original[['name', 'version', 'github_url']])
    assert actual.serial_number.tolist() == ['000123', '123456789012345678901234', 'NA', '', '批次-A01']
    assert actual.package_manager.tolist() == ['cargo', 'npm', 'None', '', 'custom']


@pytest.mark.parametrize('field', [None, 'package_manager', 'serial_number'])
def test_reader_supports_missing_optional_columns(field):
    row = dict(name='demo', version='1.0', github_url='https://github.com/example/demo')
    if field:
        row[field] = '0001'
    content = excel_bytes([row])
    actual = read_input_excel(io.BytesIO(content))
    metadata = get_input_metadata(actual.iloc[0])
    assert metadata == {key: '0001' if key == field else None
                        for key in ('package_manager', 'serial_number')}


def install_handlers(monkeypatch, module):
    completed = []

    async def github(client, url, version, name=None):
        if name == 'slow':
            await asyncio.sleep(0.04)
        completed.append(name)
        if name in ('error', 'fallback'):
            raise ValueError('controlled test failure')
        return dict(status='skipped' if name == 'skip' else 'success', license_type='MIT')

    async def archive(url, version, name, **kwargs):
        return dict(status='success', license_type='Apache-2.0')

    monkeypatch.setattr(module, 'process_github_repository', github)
    monkeypatch.setattr(module, 'process_direct_archive_url', archive)
    return completed


def assert_output(content):
    with io.BytesIO(content) as stream:
        wb = load_workbook(stream)
        sheet = wb['分析结果']
        assert [c.value for c in sheet[1]][:2] == ['serial_number', 'package_manager']
        assert [[sheet.cell(r, c).value for c in (1, 2)] for r in range(2, 7)] == [
            ['000123', 'cargo'], ['123456789012345678901234', 'npm'],
            ['NA', 'None'], [None, None], ['批次-A01', 'custom'],
        ]
        assert sheet['A2'].data_type == sheet['A3'].data_type == 's'
        headers = [c.value for c in sheet[1]]
        assert [sheet.cell(r, headers.index('status') + 1).value for r in range(2, 7)] == [
            'success', 'success', 'error', 'skipped', 'success']
        wb.close()


@pytest.mark.asyncio
async def test_cli_real_input_to_excel(monkeypatch, tmp_path):
    import main
    completed = install_handlers(monkeypatch, main)

    async def initialize():
        return object()

    monkeypatch.setattr(main, 'initialize_api', initialize)
    monkeypatch.chdir(tmp_path)
    (tmp_path / 'input.xlsx').write_bytes(excel_bytes(input_rows()))
    await main.main_async()
    assert completed.index('fast') < completed.index('slow')
    assert_output((tmp_path / 'outputs/output_latest.xlsx').read_bytes())


@pytest.mark.parametrize('route', [
    'analyze', 'analyze-and-download', 'analyze-stream',
    'analyze-stream-email', 'analyze-stream-download',
])
def test_all_api_upload_routes(monkeypatch, route):
    completed = install_handlers(monkeypatch, api)
    captured = []
    emailed = []
    original_process = api._process_repositories

    async def initialize():
        return object()

    async def process(*args, **kwargs):
        results = await original_process(*args, **kwargs)
        assert [r['serial_number'] for r in results] == [
            '000123', '123456789012345678901234', 'NA', '', '批次-A01']
        assert [r['package_manager'] for r in results] == ['cargo', 'npm', 'None', '', 'custom']
        captured.append(results)
        return results

    def email(**kwargs):
        from pathlib import Path
        assert_output(Path(kwargs['output_file_path']).read_bytes())
        emailed.append(True)
        return True

    monkeypatch.setattr(api, 'get_github_api', initialize)
    monkeypatch.setattr(api, '_process_repositories', process)
    monkeypatch.setattr(api, 'send_analysis_result', email)
    with TestClient(api.app) as client:
        response = client.post('/api/v1/' + route,
            files={'file': ('input.xlsx', excel_bytes(input_rows()),
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
            data={'email': 'test@example.com'})
    assert response.status_code == 200
    assert len(captured) == 1
    assert len(emailed) == (1 if route in ('analyze', 'analyze-stream-email') else 0)
    assert completed.index('fast') < completed.index('slow')
    if route.endswith('download'):
        assert_output(response.content)
    elif 'stream' in route:
        assert '[SUCCESS]' in response.text
        assert '无法读取Excel文件' not in response.text
    else:
        assert response.json()['status'] == 'success'


@pytest.mark.asyncio
async def test_fonts_keep_metadata_on_success_and_error(monkeypatch, tmp_path):
    import main
    import core.font_utils

    async def font(client, name, version, url):
        if name == 'error':
            raise ValueError('controlled test failure')
        return dict(status='success')

    monkeypatch.setattr(core.font_utils, 'process_font_entry', font)
    monkeypatch.chdir(tmp_path)
    results = await main.process_all_fonts(None, pd.DataFrame(input_rows()))
    assert [r['serial_number'] for r in results] == [
        '000123', '123456789012345678901234', 'NA', None, '批次-A01']
    assert results[2]['status'] == 'error'


def test_output_includes_blank_metadata_for_old_results():
    result = api._generate_output([dict(input_name='old', status='error')])
    assert result.columns.tolist()[:2] == ['serial_number', 'package_manager']
    assert result[['serial_number', 'package_manager']].isna().all().all()
