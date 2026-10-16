import pandas as pd
import pytest
from openpyxl import load_workbook

from api import _generate_output
from core.utils import format_concluded_license_semicolon, get_output_risk_levels


@pytest.mark.parametrize(("expression", "expected"), [
    ("MIT AND Apache-2.0", "MIT; Apache-2.0"),
    ("MIT AND Others", "MIT AND Others"),
    ("MIT AND Apache-2.0 AND Others", "MIT; Apache-2.0 AND Others"),
    ("(MIT OR GPL-3.0) AND LGPL-2.1", "MIT OR GPL-3.0; LGPL-2.1"),
    ("MIT OR GPL-3.0 AND LGPL-2.1", "MIT OR GPL-3.0; LGPL-2.1"),
    ("MIT and others", "MIT and others"),
    ("Brand-AND-Identifier", "Brand-AND-Identifier"),
    ("(MIT)AND(GPL-3.0)", "MIT; GPL-3.0"),
])
def test_semicolon_column_formats_license_items(expression, expected):
    assert format_concluded_license_semicolon(expression) == expected


@pytest.mark.parametrize(("expression", "expected"), [
    ("MIT", "低风险"),
    ("MIT AND GPL-3.0-only", "低风险;高风险"),
    ("MIT OR GPL-3.0-only", "低风险"),
    ("(MIT OR GPL-3.0-only) AND LGPL-2.1", "低风险;中风险"),
    ("MIT OR GPL-3.0-only AND LGPL-2.1", "低风险;中风险"),
    ("GPL-3.0-only OR (MIT AND LGPL-2.1)", "低风险;中风险"),
    ("(MIT AND GPL-3.0-only) OR (Apache-2.0 AND MPL-2.0)", "低风险;低风险;中风险"),
    ("GPL-2.0-only WITH Classpath-exception-2.0 AND MIT", "高风险;低风险"),
    ("MIT AND Others", "低风险"),
    ("MIT AND SomeWeirdLicense", "低风险;未知风险"),
    ("MIT OR SomeWeirdLicense", "低风险"),
    ("Others", "未知风险"),
    ("Unlicensed", "高风险"),
])
def test_output_risk_lists_selected_licenses_in_and_order(expression, expected):
    assert get_output_risk_levels(expression) == expected


def test_api_output_places_new_column_after_original_expression():
    output = _generate_output([
        {"input_name": "demo", "license_type": "MIT AND GPL-3.0-only", "status": "success"},
        {"input_name": "other", "license_type": "MIT", "status": "success",
         "license_analysis": {"thirdparty_dirs": ["vendor"]}},
    ])

    assert output.columns.get_loc("concluded_license_semicolon") == output.columns.get_loc("concluded_license") + 1
    assert output["concluded_license"].tolist() == ["MIT AND GPL-3.0-only", "MIT AND Others"]
    assert output["concluded_license_semicolon"].tolist() == ["MIT; GPL-3.0-only", "MIT AND Others"]
    assert output["risk_level"].tolist() == ["低风险;高风险", "低风险"]


@pytest.mark.asyncio
async def test_cli_excel_has_new_column_and_separate_risks(monkeypatch, tmp_path):
    import main

    async def fake_initialize_api():
        return object()

    async def fake_process_all_repos(api, df, max_concurrency):
        return [{"input_name": "demo", "license_type": "(MIT OR GPL-3.0-only) AND LGPL-2.1",
                 "status": "success"}]

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(main, "initialize_api", fake_initialize_api)
    monkeypatch.setattr(main, "process_all_repos", fake_process_all_repos)
    monkeypatch.setattr(main, "read_input_excel", lambda path: pd.DataFrame([{}]))

    await main.main_async()

    workbook = load_workbook(tmp_path / "outputs" / "output_latest.xlsx", read_only=True)
    rows = list(workbook["分析结果"].values)
    headers = rows[0]
    assert headers.index("concluded_license_semicolon") == headers.index("concluded_license") + 1
    assert rows[1][headers.index("concluded_license")] == "(MIT OR GPL-3.0-only) AND LGPL-2.1"
    assert rows[1][headers.index("concluded_license_semicolon")] == "MIT OR GPL-3.0-only; LGPL-2.1"
    assert rows[1][headers.index("risk_level")] == "低风险;中风险"


DEDUPLICATED_CASES = [
    ('MIT AND (Apache-2.0 AND MIT)', 'MIT; Apache-2.0', '低风险;低风险'),
    ('(MIT OR Apache-2.0) AND (MIT OR Apache-2.0)', 'MIT OR Apache-2.0', '低风险'),
    ('MIT AND MIT AND Others', 'MIT AND Others', '低风险'),
    ('MIT AND Apache-2.0 AND MIT AND Others', 'MIT; Apache-2.0 AND Others', '低风险;低风险'),
    ('((MIT AND MIT) AND (GPL-3.0-only AND MIT))', 'MIT; GPL-3.0-only', '低风险;高风险'),
    ('MIT AND (MIT OR GPL-3.0-only)', 'MIT; MIT OR GPL-3.0-only', '低风险;低风险'),
    ('GPL-3.0-only OR (MIT AND LGPL-2.1)', 'GPL-3.0-only OR MIT; LGPL-2.1', '低风险;中风险'),
    ('(MIT AND GPL-3.0-only) OR (Apache-2.0 AND MPL-2.0)',
     'MIT; GPL-3.0-only OR Apache-2.0; MPL-2.0', '低风险;低风险;中风险'),
    ('GPL-2.0-only WITH Classpath-exception-2.0 AND (GPL-2.0-only WITH Classpath-exception-2.0)',
     'GPL-2.0-only WITH Classpath-exception-2.0', '高风险'),
    ('GPL-2.0-only WITH Classpath-exception-2.0 AND GPL-2.0-only',
     'GPL-2.0-only WITH Classpath-exception-2.0; GPL-2.0-only', '高风险;高风险'),
    ('GPL-3.0-only AND GPL-3.0-or-later', 'GPL-3.0-only; GPL-3.0-or-later', '高风险;高风险'),
    ('  MIT and (mit) AND Apache-2.0  and others ', 'MIT; Apache-2.0 and others', '低风险;低风险'),
    ('MIT; ; MIT; (Apache-2.0);', 'MIT; Apache-2.0', '低风险;低风险'),
    ('SomeWeirdLicense AND SomeWeirdLicense AND MIT', 'SomeWeirdLicense; MIT', '未知风险;低风险'),
    ('Brand-AND-Identifier AND Brand-AND-Identifier', 'Brand-AND-Identifier', '未知风险'),
    ('MIT OR AND Apache-2.0', 'MIT OR; Apache-2.0', '低风险;低风险'),
]


@pytest.mark.parametrize(('expression', 'display', 'risks'), DEDUPLICATED_CASES)
def test_display_and_risks_share_deduplicated_items(expression, display, risks):
    assert format_concluded_license_semicolon(expression) == display
    assert get_output_risk_levels(expression) == risks
    assert get_output_risk_levels(display) == risks
    assert len(display.split(';')) == len(risks.split(';'))
    assert '(' not in display and ')' not in display


@pytest.mark.parametrize('value', [None, float('nan'), '', '   '])
def test_blank_expression_has_no_license_items(value):
    display = format_concluded_license_semicolon(value)
    assert display is None or pd.isna(display) or display == ''
    assert get_output_risk_levels(value) == '未知风险'


def test_api_dataframe_keeps_original_expression_and_aligns_deduplicated_risk():
    cases = DEDUPLICATED_CASES[:15]
    output = _generate_output([
        {'input_name': f'demo{index}', 'license_type': expression, 'status': 'success'}
        for index, (expression, _, _) in enumerate(cases)
    ])
    assert output.concluded_license.tolist() == [expression.strip() for expression, _, _ in cases]
    assert output.concluded_license_semicolon.tolist() == [display for _, display, _ in cases]
    assert output.risk_level.tolist() == [risks for _, _, risks in cases]


@pytest.mark.parametrize('entry', ['cli', 'api'])
def test_actual_excel_export_has_deduplicated_licenses_and_matching_risks(monkeypatch, tmp_path, entry):
    import asyncio
    import io
    from unittest.mock import AsyncMock
    from fastapi.testclient import TestClient
    import api
    import main

    cases = DEDUPLICATED_CASES[:15]
    results = [dict(input_name=f'demo{i}', license_type=expression, status='success')
               for i, (expression, _, _) in enumerate(cases)]
    source = io.BytesIO()
    pd.DataFrame([{'name': r['input_name'], 'github_url': 'https://github.com/example/demo'}
                  for r in results]).to_excel(source, index=False)
    monkeypatch.chdir(tmp_path)
    if entry == 'cli':
        (tmp_path / 'input.xlsx').write_bytes(source.getvalue())
        monkeypatch.setattr(main, 'initialize_api', AsyncMock(return_value=object()))
        monkeypatch.setattr(main, 'process_all_repos', AsyncMock(return_value=results))
        asyncio.run(main.main_async())
        content = (tmp_path / 'outputs/output_latest.xlsx').read_bytes()
    else:
        monkeypatch.setattr(api, 'get_github_api', AsyncMock(return_value=object()))
        monkeypatch.setattr(api, '_process_repositories', AsyncMock(return_value=results))
        with TestClient(api.app) as client:
            response = client.post('/api/v1/analyze-and-download', files={
                'file': ('input.xlsx', source.getvalue(),
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert response.status_code == 200
        content = response.content
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(workbook['分析结果'].values)
    for values, (expression, display, risks) in zip(rows[1:], cases, strict=True):
        record = dict(zip(rows[0], values))
        assert record['concluded_license'] == expression.strip()
        assert record['concluded_license_semicolon'] == display
        assert record['risk_level'] == risks
    workbook.close()
